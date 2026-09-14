#!/usr/bin/env python3
"""Independent read-only OOXML / openpyxl release gate; never rewrites XLSX."""
import sys,zipfile,xml.etree.ElementTree as ET,json,warnings,sqlite3
from pathlib import Path
from openpyxl import load_workbook
file=Path(sys.argv[1]);dbfile=Path(sys.argv[2]) if len(sys.argv)>2 else None
expected=['Workbook Guide','Executive Dashboard','Student Master','Student 360','Parent Directory','Academic Profiles','Class Register','Class Hours Summary','Monthly Class Hours','Attendance Register','Attendance Summary','Monthly Attendance','Fee Register','Payments Ledger','Student Finance','Outstanding Dues','Monthly Collections','Payment Methods','Test Register','Performance Summary','Subject Performance','Intake Submissions','Admissions Pipeline','Data Quality','Audit Log','System Settings','Data Dictionary','Parent Accounts','Parent Student Links','Enrollments','Course History','Old Records Alumni','Payment Proofs','Historical Payments','Renewal Requests','Portal Activity','Teachers','Batches','Subjects','Assignments','Announcements','Users & Roles','Backup History','Notification Summary']
checks=0
with zipfile.ZipFile(file) as z:
 assert z.testzip() is None;checks+=1
 xmls=[n for n in z.namelist() if n.endswith(('.xml','.rels'))]
 for n in xmls:ET.fromstring(z.read(n));checks+=1
 assert not any(n.startswith('xl/media/') for n in z.namelist());checks+=1
with warnings.catch_warnings(record=True) as caught:
 warnings.simplefilter('always');wb=load_workbook(file,data_only=False)
 assert not caught,[str(w.message) for w in caught];checks+=1
assert wb.sheetnames==expected,wb.sheetnames;checks+=1
bad={'#REF!','#DIV/0!','#VALUE!','#NAME?','#N/A','#NUM!','#NULL!'}
for ws in wb:
 assert ws.freeze_panes,ws.title+' missing freeze panes';checks+=1
 for row in ws:
  for c in row:
   assert c.data_type!='e' and c.value not in bad if isinstance(c.value,str) else c.data_type!='e',(ws.title,c.coordinate,c.value)
   if c.has_style:assert c.style_id<len(wb._cell_styles)
 ranges=list(ws.merged_cells.ranges)
 for i,a in enumerate(ranges):
  for b in ranges[i+1:]:assert a.max_row<b.min_row or b.max_row<a.min_row or a.max_col<b.min_col or b.max_col<a.min_col,(ws.title,a,b)
 checks+=1
assert wb['Workbook Guide']['D3'].value=='7.0' and wb['Workbook Guide']['F3'].value==44;checks+=1
# The saved fixture uses adversarial but artificial history/current records.
if dbfile:
 conn=sqlite3.connect(dbfile);conn.row_factory=sqlite3.Row
 def values(sheet):return [[c.value for c in row] for row in wb[sheet]]
 assert not any('QA Returning Alumnus' in str(v) for row in values('Old Records Alumni') for v in row);checks+=1
 assert any('QA Returning Alumnus' in str(v) for row in values('Course History') for v in row);checks+=1
 assert any('QA Excel failure survivor' in str(v) for row in values('Student Master') for v in row);checks+=1
 # No credential field or token exported in account worksheets.
 for sheet in ['Users & Roles','Parent Accounts']:
  assert not any(any(secret in str(v).lower() for secret in ['password_hash','token_hash','session_secret']) for row in values(sheet) for v in row);checks+=1
 conn.close()
report={'passed':True,'checks':checks,'parser':'openpyxl','worksheets':wb.sheetnames,'xml_parts':len(xmls),'zip_integrity':'ok','formula_errors':0,'parser_warnings':0,'bytes':file.stat().st_size}
print(json.dumps(report,indent=2));file.with_suffix('.validation.json').write_text(json.dumps(report,indent=2))
